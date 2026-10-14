"""Робить файл для Kahoot із запитань вікторини (dani.js → kviz).
Запуск: python3 _robota/gutirky/budova-la-1/kahoot.py
Розкладка — як в офіційному шаблоні Kahoot: заголовки в рядку 8, запитання з рядка 9,
стовпець A — номер, B — запитання (≤120), C–F — відповіді (≤75), G — час, H — номер правильної."""
import json, re, subprocess, pathlib
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill

TUT = pathlib.Path(__file__).parent
js = (TUT / 'dani.js').read_text(encoding='utf-8')
# беремо дані через node, щоб не розбирати JS руками
kviz = json.loads(subprocess.check_output(['node', '-e', f"global.window={{}};{js};console.log(JSON.stringify(window.LA.kviz))"]))

wb = Workbook(); ws = wb.active; ws.title = 'Kahoot'
F = 'Arial'
ws['B2'] = 'Будова літальних апаратів І — вікторина для Kahoot (15 запитань)'; ws['B2'].font = Font(name=F, size=14, bold=True)
ws['B3'] = 'Kahoot: Створити → Додати запитання → Імпорт → Імпортувати таблицю → вибрати цей файл.'; ws['B3'].font = Font(name=F, size=10)
ws['B4'] = 'Правильна відповідь — номер варіанта (1–4). Час на кожне запитання — 20 секунд.'; ws['B4'].font = Font(name=F, size=10)
ws['B5'] = 'Запитання відповідають вимогам вмілості «Будова літальних апаратів І» (Летунська, проба 1).'; ws['B5'].font = Font(name=F, size=10)
H = ['', 'Question - max 120 characters', 'Answer 1 - max 75 characters', 'Answer 2 - max 75 characters',
     'Answer 3 - max 75 characters', 'Answer 4 - max 75 characters',
     'Time limit (sec) – 5, 10, 20, 30, 60, 90, 120, or 240 secs', 'Correct answer(s) - choose at least one']
fill = PatternFill('solid', start_color='46178F')
for c, h in enumerate(H, 1):
    cell = ws.cell(row=8, column=c, value=h or None)
    cell.font = Font(name=F, bold=True, color='FFFFFF'); cell.fill = fill
    cell.alignment = Alignment(wrap_text=True, vertical='center')
ws.row_dimensions[8].height = 44
probl = []
for i, q in enumerate(kviz, 1):
    r = 8 + i
    if len(q['q']) > 120: probl.append(f"запитання {i}: {len(q['q'])} знаків")
    for a in q['a']:
        if len(a) > 75: probl.append(f"відповідь у {i}: {len(a)} знаків")
    row = [i, q['q'], *q['a'], 20, q['ok'] + 1]
    for c, v in enumerate(row, 1):
        cell = ws.cell(row=r, column=c, value=v)
        cell.font = Font(name=F, size=11); cell.alignment = Alignment(wrap_text=True, vertical='top')
for col, w in zip('ABCDEFGH', [5, 52, 26, 26, 26, 26, 14, 14]):
    ws.column_dimensions[col].width = w
ws.freeze_panes = 'B9'
out = TUT / 'budova-la-1-kahoot.xlsx'
wb.save(out)
print('збережено', out, '— запитань:', len(kviz), '| перевищення:', probl or 'немає')
